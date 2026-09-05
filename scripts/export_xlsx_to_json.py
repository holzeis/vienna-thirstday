#!/usr/bin/env python3
"""
Exports one season's "Kicken_<year>.xlsx" spreadsheet (its "Ergebnisse <year>"
sheet) into a JSON file the backend's historical-data importer can load. Each
season lives in its own workbook/file; run this once per year.

The original spreadsheet recorded, per player per matchday, the exact POINTS
(4 win / 2 draw / 1 loss) and GOAL DIFFERENCE for that player. We preserve
those two numbers per player exactly (no recomputation), so the season
standings reconstructed from this import match the original spreadsheet's
table exactly. We additionally group players into team A (points == 4, or the
first half of a draw) / team B (points == 1, or the second half of a draw)
purely for display purposes - the legacy sheet never recorded actual team
rosters or the literal final score, so those are best-effort reconstructions
and are called out as such in the README.

Usage:
    python3 export_xlsx_to_json.py /path/to/Kicken_2026.xlsx /path/to/output.json [sheet_name]

`sheet_name` is optional - if omitted, the script picks the sole sheet named
"Ergebnisse <year>" (erroring out if there isn't exactly one match).
"""
import json
import re
import sys
from collections import Counter
from datetime import datetime, timezone

import openpyxl


def find_results_sheet(wb, explicit_name):
    if explicit_name:
        return explicit_name
    candidates = [name for name in wb.sheetnames if re.fullmatch(r"Ergebnisse \d{4}", name)]
    if len(candidates) == 1:
        return candidates[0]
    raise SystemExit(
        f"Could not auto-detect the results sheet (found {candidates or 'none'} matching "
        f"'Ergebnisse <year>' among {wb.sheetnames}). Pass the sheet name explicitly as the 3rd argument."
    )


def main():
    if len(sys.argv) not in (3, 4):
        print(__doc__)
        sys.exit(1)

    xlsx_path, out_path = sys.argv[1], sys.argv[2]
    explicit_sheet = sys.argv[3] if len(sys.argv) == 4 else None
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[find_results_sheet(wb, explicit_sheet)]

    rows = list(ws.iter_rows(values_only=True))
    name_row = rows[2]  # row index 2 (0-based) has player names in the "P" column of each pair
    header_row = rows[3]  # 'Spieltag','Datum','P','T','P2','T2', ...

    # Map each column index whose header starts with 'P' (points) to the player name
    # found in the same column of name_row, and record the paired goal-diff column.
    player_columns = []  # list of (points_col_idx, goaldiff_col_idx, player_name)
    for col_idx, header in enumerate(header_row):
        if header is None:
            continue
        header_str = str(header)
        if header_str.startswith("P") and not header_str.startswith("Punkte"):
            name = name_row[col_idx]
            if name:
                player_columns.append((col_idx, col_idx + 1, str(name).strip()))

    players = sorted({name for _, _, name in player_columns})

    gamedays = []
    for row in rows[4:]:
        spieltag = row[0]
        raw_date = row[1]
        if spieltag is None or raw_date is None:
            continue
        if isinstance(raw_date, datetime):
            date_str = raw_date.date().isoformat()
        else:
            date_str = str(raw_date)

        entries = []  # {name, points, goalDiff}
        for points_col, goaldiff_col, name in player_columns:
            points = row[points_col] if points_col < len(row) else None
            gd = row[goaldiff_col] if goaldiff_col < len(row) else None
            if points is None:
                continue
            entries.append({"name": name, "points": int(points), "goalDiff": int(gd or 0)})

        if not entries:
            continue

        winners = [e for e in entries if e["points"] == 4]
        losers = [e for e in entries if e["points"] == 1]
        drawers = [e for e in entries if e["points"] not in (4, 1)]

        if winners or losers:
            team_a = winners
            team_b = losers
        else:
            half = len(drawers) // 2
            team_a = drawers[:half]
            team_b = drawers[half:]

        # Best-effort placeholder score for display only (standings use the
        # exact per-player points/goalDiff above, not this reconstructed score).
        # Every player on a team should share the same team-level goal
        # difference, so cross-check both sides' recorded values instead of
        # only looking at the winners - a lone mistyped cell on one side
        # shouldn't silently win out over an agreeing majority on the other.
        candidates = [abs(e["goalDiff"]) for e in winners if e["goalDiff"] != 0]
        candidates += [abs(e["goalDiff"]) for e in losers if e["goalDiff"] != 0]
        if candidates:
            counts = Counter(candidates)
            best_count = max(counts.values())
            most_common = [v for v, c in counts.items() if c == best_count]
            if len(most_common) > 1:
                winner_values = {abs(e["goalDiff"]) for e in winners if e["goalDiff"] != 0}
                preferred = [v for v in most_common if v in winner_values]
                placeholder_diff = min(preferred) if preferred else min(most_common)
            else:
                placeholder_diff = most_common[0]
        else:
            placeholder_diff = 0

        gamedays.append(
            {
                "spieltag": spieltag,
                "date": date_str,
                "teamA": [{"name": e["name"], "points": e["points"], "goalDiff": e["goalDiff"]} for e in team_a],
                "teamB": [{"name": e["name"], "points": e["points"], "goalDiff": e["goalDiff"]} for e in team_b],
                "placeholderScoreA": placeholder_diff,
                "placeholderScoreB": 0,
            }
        )

    output = {
        "source": xlsx_path,
        "exportedAt": datetime.now(timezone.utc).isoformat(),
        "players": players,
        "gamedays": gamedays,
    }

    with open(out_path, "w") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)

    print(f"Wrote {len(players)} players and {len(gamedays)} played gamedays to {out_path}")


if __name__ == "__main__":
    main()
